"""Tests for Phase 2 resume parsing (PDF/DOCX extraction)."""

from io import BytesIO

import fitz  # PyMuPDF
import pytest
from docx import Document

from app.services.resume_parser import (
    ResumeDocument,
    parse_docx,
    parse_pdf,
    parse_resume,
)


def make_pdf_bytes(text: str) -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), text)
    buf = BytesIO()
    doc.save(buf)
    doc.close()
    return buf.getvalue()


def make_docx_bytes(paragraphs: list[str]) -> bytes:
    document = Document()
    for para in paragraphs:
        document.add_paragraph(para)
    buf = BytesIO()
    document.save(buf)
    return buf.getvalue()


def test_valid_pdf_extraction():
    pdf_bytes = make_pdf_bytes("Jane Doe\nSoftware Engineer with Python experience")
    doc = parse_resume(pdf_bytes, "resume.pdf")
    assert isinstance(doc, ResumeDocument)
    assert doc.filename == "resume.pdf"
    assert doc.file_type == "pdf"
    assert "Jane Doe" in doc.text
    assert "Software Engineer" in doc.text
    assert doc.file_size == len(pdf_bytes)
    assert doc.character_count == len(doc.text)
    assert doc.word_count == len(doc.text.split())


def test_valid_docx_extraction():
    docx_bytes = make_docx_bytes(["John Smith", "Data Analyst skilled in SQL"])
    doc = parse_resume(docx_bytes, "resume.docx")
    assert isinstance(doc, ResumeDocument)
    assert doc.filename == "resume.docx"
    assert doc.file_type == "docx"
    assert "John Smith" in doc.text
    assert "Data Analyst" in doc.text
    assert doc.file_size == len(docx_bytes)


def test_uppercase_extension_accepted():
    pdf_bytes = make_pdf_bytes("Uppercase extension content")
    doc = parse_resume(pdf_bytes, "RESUME.PDF")
    assert doc.file_type == "pdf"
    assert "Uppercase" in doc.text


def test_unsupported_file_type():
    with pytest.raises(ValueError, match="Unsupported file type"):
        parse_resume(b"some text", "resume.txt")


def test_missing_extension_rejected():
    with pytest.raises(ValueError, match="Unsupported file type"):
        parse_resume(b"some text", "resume")


def test_empty_file_rejected():
    with pytest.raises(ValueError, match="empty"):
        parse_resume(b"", "resume.pdf")


def test_corrupted_pdf_rejected():
    with pytest.raises(ValueError, match="Could not parse PDF"):
        parse_resume(b"this is not a pdf at all %%%", "broken.pdf")


def test_corrupted_docx_rejected():
    with pytest.raises(ValueError, match="Could not parse DOCX"):
        parse_resume(b"this is not a docx file", "broken.docx")


def test_pdf_with_no_extractable_text_rejected():
    doc = fitz.open()
    doc.new_page()  # blank page, no text
    buf = BytesIO()
    doc.save(buf)
    doc.close()
    with pytest.raises(ValueError, match="No extractable text"):
        parse_pdf(buf.getvalue())


def test_docx_with_no_extractable_text_rejected():
    document = Document()  # no paragraphs
    buf = BytesIO()
    document.save(buf)
    with pytest.raises(ValueError, match="No extractable text"):
        parse_docx(buf.getvalue())


def test_metadata_calculation_and_whitespace_normalization():
    docx_bytes = make_docx_bytes(["  Hello    world  ", "", "   Python   developer   "])
    doc = parse_resume(docx_bytes, "test.docx")
    assert "  " not in doc.text.replace("\n\n", "\n")
    assert doc.character_count == len(doc.text)
    assert doc.word_count == len(doc.text.split())
    assert doc.word_count == 4  # Hello world Python developer
