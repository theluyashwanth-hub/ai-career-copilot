"""Resume upload and text-extraction backend (Phase 2).

Supports PDF (PyMuPDF) and DOCX (python-docx) only.
No LLM, RAG, database, or agent logic lives here.
"""

import re
from io import BytesIO

import fitz  # PyMuPDF
from docx import Document
from pydantic import BaseModel, Field

SUPPORTED_EXTENSIONS = (".pdf", ".docx")


class ResumeDocument(BaseModel):
    """Validated resume extraction result."""

    filename: str = Field(min_length=1)
    file_type: str
    file_size: int = Field(ge=0)
    text: str = Field(min_length=1)
    character_count: int = Field(ge=0)
    word_count: int = Field(ge=0)


def _normalize_text(raw_text: str) -> str:
    """Collapse unnecessary whitespace while keeping text readable."""
    text = raw_text.replace("\r\n", "\n").replace("\r", "\n")
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.split("\n")]
    # Drop leading/trailing blank lines but keep single blank lines inside.
    cleaned: list[str] = []
    for line in lines:
        if line == "" and (not cleaned or cleaned[-1] == ""):
            continue
        cleaned.append(line)
    while cleaned and cleaned[0] == "":
        cleaned.pop(0)
    while cleaned and cleaned[-1] == "":
        cleaned.pop()
    return "\n".join(cleaned).strip()


def _build_document(filename: str, file_type: str, file_size: int, raw_text: str) -> ResumeDocument:
    text = _normalize_text(raw_text)
    if not text:
        raise ValueError(f"No extractable text found in '{filename}'.")
    return ResumeDocument(
        filename=filename,
        file_type=file_type,
        file_size=file_size,
        text=text,
        character_count=len(text),
        word_count=len(text.split()),
    )


def parse_pdf(file_bytes: bytes) -> str:
    """Extract and normalize text from PDF bytes."""
    if not file_bytes:
        raise ValueError("Uploaded file is empty.")
    try:
        with fitz.open(stream=file_bytes, filetype="pdf") as doc:
            parts = [page.get_text() or "" for page in doc]
    except Exception as exc:
        raise ValueError(f"Could not parse PDF file: {exc}") from exc
    text = _normalize_text("\n".join(parts))
    if not text:
        raise ValueError("No extractable text found in PDF file.")
    return text


def parse_docx(file_bytes: bytes) -> str:
    """Extract and normalize text from DOCX bytes."""
    if not file_bytes:
        raise ValueError("Uploaded file is empty.")
    try:
        document = Document(BytesIO(file_bytes))
        parts: list[str] = [p.text or "" for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text or "")
    except Exception as exc:
        raise ValueError(f"Could not parse DOCX file: {exc}") from exc
    text = _normalize_text("\n".join(parts))
    if not text:
        raise ValueError("No extractable text found in DOCX file.")
    return text


def parse_resume(file_bytes: bytes, filename: str) -> ResumeDocument:
    """Validate, extract, and build a ResumeDocument from uploaded bytes."""
    if not filename or "." not in filename:
        raise ValueError(f"Unsupported file type: '{filename}'. Only PDF and DOCX are supported.")
    if not file_bytes:
        raise ValueError(f"Uploaded file '{filename}' is empty.")

    extension = "." + filename.rsplit(".", 1)[-1].lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type: '{extension}'. Only PDF (.pdf) and DOCX (.docx) are supported."
        )

    file_size = len(file_bytes)
    if extension == ".pdf":
        raw_text = parse_pdf(file_bytes)
        file_type = "pdf"
    else:
        raw_text = parse_docx(file_bytes)
        file_type = "docx"

    return _build_document(
        filename=filename,
        file_type=file_type,
        file_size=file_size,
        raw_text=raw_text,
    )
